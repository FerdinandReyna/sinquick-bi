# -*- coding: utf-8 -*-
"""Llena las columnas "Resultado obtenido" y "Estado" del plan de pruebas con los
resultados reales de una ejecución del pipeline.

Uso (desde la raíz del repositorio, con la carpeta reportes/ descargada de Jenkins):
    python docs/plan-pruebas/completar_resultados.py [carpeta_reportes]

Lee:  reportes/tsqlt.xml            (pruebas unitarias, formato JUnit)
      reportes/resultados-pytest.json (pruebas de pytest, generado por conftest.py)
Escribe: docs/plan-pruebas/Plan_de_Pruebas_SIN_Quick_ejecutado.docx
Requiere: pip install python-docx
"""
import json
import re
import sys
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import RGBColor

AQUI = Path(__file__).resolve().parent
REPORTES = Path(sys.argv[1]) if len(sys.argv) > 1 else AQUI.parents[1] / "reportes"
ENTRADA = AQUI / "Plan_de_Pruebas_SIN_Quick.docx"
SALIDA = AQUI / "Plan_de_Pruebas_SIN_Quick_ejecutado.docx"


def resultados_tsqlt(ruta):
    res = {}
    if not ruta.exists():
        return res
    for caso in ET.parse(ruta).getroot().iter("testcase"):
        m = re.search(r"CP-\d+", caso.get("name", ""))
        if not m:
            continue
        falla = caso.find("failure") if caso.find("failure") is not None else caso.find("error")
        tiempo = caso.get("time", "0")
        if falla is None:
            res[m.group()] = ("APROBADO", f"Prueba tSQLt aprobada en {float(tiempo):.2f} s.")
        else:
            msg = (falla.get("message") or falla.text or "").strip()
            res[m.group()] = ("FALLIDO", f"Prueba tSQLt fallida: {msg[:250]}")
    return res


def resultados_pytest(ruta):
    res = {}
    if not ruta.exists():
        return res
    for r in json.loads(ruta.read_text(encoding="utf-8")):
        texto = r.get("obtenido") or ""
        if r["estado"] == "FALLIDO":
            texto = f"Falló: {r.get('mensaje', '')}"
        elif r["estado"] == "OMITIDO":
            texto = f"No ejecutado: {r.get('mensaje', '')}"
        res[r["caso"]] = (r["estado"], f"{texto} ({r['duracion_s']} s)")
    return res


def escribir(celda, texto, color):
    celda.text = ""
    run = celda.paragraphs[0].add_run(texto)
    run.font.color.rgb = color


def main():
    res = {**resultados_tsqlt(REPORTES / "tsqlt.xml"), **resultados_pytest(REPORTES / "resultados-pytest.json")}
    if not res:
        sys.exit(f"No se encontraron resultados en {REPORTES}")
    doc = Document(ENTRADA)
    llenados = 0
    for tabla in doc.tables:
        if tabla.rows[0].cells[0].text.strip() != "ID":
            continue
        caso = tabla.rows[0].cells[1].text.strip()
        if caso not in res:
            continue
        estado, texto = res[caso]
        color = RGBColor(0x1E, 0x7B, 0x34) if estado == "APROBADO" else RGBColor(0xB4, 0x23, 0x18)
        for fila in tabla.rows:
            clave = fila.cells[0].text.strip()
            if clave == "Resultado obtenido":
                escribir(fila.cells[1], texto, RGBColor(0x13, 0x29, 0x3D))
            elif clave == "Estado":
                escribir(fila.cells[1], estado.capitalize(), color)
        llenados += 1
    doc.core_properties.comments = f"Resultados registrados el {datetime.now():%Y-%m-%d %H:%M}"
    doc.save(SALIDA)
    aprobados = sum(1 for e, _ in res.values() if e == "APROBADO")
    print(f"{llenados} casos completados ({aprobados} aprobados, {len(res) - aprobados} no aprobados).")
    print(f"Guardado en {SALIDA}")


if __name__ == "__main__":
    main()
